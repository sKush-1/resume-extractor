"""
DOCX text extractor using python-docx.
Extracts text from Word document paragraphs.
"""

import io
from docx import Document

from ..logger import create_logger

logger = create_logger("docx-extractor")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text content from a DOCX file.

    Args:
        file_bytes: Raw bytes of the DOCX file

    Returns:
        Extracted text string

    Raises:
        ValueError: If the DOCX is empty or unreadable
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []
        
        # Extract from paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text)
                
        # Extract from tables (very common in resume headers)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            text_parts.append(p.text)

        full_text = "\n".join(text_parts).strip()

        if not full_text:
            raise ValueError("DOCX contains no extractable text")

        logger.debug("DOCX text extracted", extra={"length": len(full_text)})
        return full_text

    except Exception as e:
        if "not extractable" in str(e):
            raise
        raise ValueError(f"Invalid or corrupted DOCX: {e}") from e
